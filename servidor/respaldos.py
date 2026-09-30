#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Respaldo de la clínica — exportar e importar datos en JSON y Excel.

Reglas del respaldo:
  - El Excel lleva una hoja por tabla, con el nombre exacto de la tabla.
  - La primera hoja es INSTRUCCIONES: explica el formato y las restricciones.
  - La importación solo AGREGA lo que falta y ACTUALIZA lo que coincide por su
    clave natural. Nunca borra registros.
  - Las tablas hijas (CITA, ATENCION, BOLETA, etc.) se reenganchan solas: el
    Id de cada registro del archivo se traduce al Id nuevo de esta base.
  - USUARIO y AUDITORIA se exportan como referencia pero no se importan: las
    contraseñas se guardan cifradas y no salen en el archivo.
"""

from __future__ import annotations

import io
from datetime import date, datetime, time
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from datos import BASE_DATOS, Tx, columnas, consultar, escalar

HOJA_INSTRUCCIONES = "INSTRUCCIONES"

# Orden obligatorio: primero los catálogos, después los movimientos que
# dependen de ellos (respeta las claves foráneas del modelo físico).
TABLAS = [
    {
        "nombre": "CLIENTE",
        "etiqueta": "Clientes",
        "detalle": "Pacientes de la clínica.",
        "restricciones": "Noms_Cliente es obligatorio. Dni_Cliente debe tener 8 dígitos y no repetirse.",
    },
    {
        "nombre": "PODOLOGO",
        "etiqueta": "Podólogos",
        "detalle": "Profesionales que ejecutan las atenciones.",
        "restricciones": "Nom_Podologo es obligatorio.",
    },
    {
        "nombre": "RECEPCIONISTA",
        "etiqueta": "Recepcionistas",
        "detalle": "Personal de counter que agenda y cobra.",
        "restricciones": "Noms_Recep es obligatorio.",
    },
    {
        "nombre": "SERVICIO",
        "etiqueta": "Servicios",
        "detalle": "Catálogo de procedimientos y sus precios.",
        "restricciones": "Nom_serv es obligatorio. Precio_Servicio no puede ser negativo.",
    },
    {
        "nombre": "MEDICAMENTO",
        "etiqueta": "Medicamentos",
        "detalle": "Inventario de productos que se entregan al paciente.",
        "restricciones": "Nom_Med es obligatorio. Stock_Med no puede ser negativo.",
    },
    {
        "nombre": "CITA",
        "etiqueta": "Citas",
        "detalle": "Agenda: une cliente, podólogo y recepcionista.",
        "restricciones": "Regla R1: no puede haber dos citas activas del mismo podólogo a la "
                         "misma fecha y hora. Estado: Programada, Atendida, Cancelada o No asistio.",
    },
    {
        "nombre": "ATENCION",
        "etiqueta": "Atenciones",
        "detalle": "Registro clínico: diagnóstico, tratamiento y observaciones.",
        "restricciones": "Una sola atención por cita. Diag y Trat son texto libre.",
    },
    {
        "nombre": "DETALLE_ATENCION",
        "etiqueta": "Detalle de servicios",
        "detalle": "Servicios aplicados en cada atención.",
        "restricciones": "Cant_Aten debe ser mayor que cero. Subt_Aten = cantidad x precio del servicio.",
    },
    {
        "nombre": "DETALLE_MEDICAMENTO",
        "etiqueta": "Detalle de medicamentos",
        "detalle": "Medicamentos entregados en cada atención.",
        "restricciones": "Cant_Med debe ser mayor que cero y descontar stock. "
                         "Subt_Med = cantidad x precio del medicamento.",
    },
    {
        "nombre": "BOLETA",
        "etiqueta": "Boletas",
        "detalle": "Comprobante de pago de una atención.",
        "restricciones": "Regla R2: Importe = suma de los detalles de la atención. "
                         "Estado_Pago: Pendiente o Pagada. Regla R4: una boleta pagada no se "
                         "modifica ni se elimina.",
    },
    {
        "nombre": "PAGO",
        "etiqueta": "Pagos",
        "detalle": "Pagos aplicados a una boleta.",
        "restricciones": "Regla R2: la suma de pagos no puede superar el importe de la boleta. "
                         "Tipo_Pago: Efectivo, Tarjeta, Yape o Transferencia.",
    },
    {
        "nombre": "USUARIO",
        "etiqueta": "Usuarios (solo referencia)",
        "detalle": "Credenciales del sistema, para consulta.",
        "restricciones": "NO se importa. Por seguridad la contraseña nunca se exporta; "
                         "los usuarios se crean en el módulo Usuarios y roles.",
        "importar": False,
        "columnas_exportar": ["Id_Usuario", "Usuario", "Rol", "Id_Podologo", "Id_Recep",
                              "Nombres", "Activo", "Fe_Alta", "Ultimo_Acceso"],
    },
    {
        "nombre": "AUDITORIA",
        "etiqueta": "Auditoría (solo referencia)",
        "detalle": "Bitácora de accesos y cambios.",
        "restricciones": "NO se importa: es la bitácora de seguridad del sistema.",
        "importar": False,
    },
]

# Clave natural con la que se reconoce un registro ya existente.
CLAVES = {
    "CLIENTE": ("Dni_Cliente",),
    "PODOLOGO": ("Email_Podo",),
    "RECEPCIONISTA": ("Tel_Recep",),
    "SERVICIO": ("Nom_serv",),
    "MEDICAMENTO": ("Nom_Med",),
    "USUARIO": ("Usuario",),
}
NOMBRE_TABLA = {t["nombre"]: t for t in TABLAS}


# ---------------------------------------------------------------------------
#  Conversión de valores
# ---------------------------------------------------------------------------
def a_texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def a_valor(valor):
    """Normaliza lo que viene de Excel o JSON antes de guardarlo."""
    if valor is None:
        return None
    if isinstance(valor, str):
        limpio = valor.strip()
        return limpio or None
    if isinstance(valor, datetime):
        return valor.date() if valor.time() == time(0, 0) else valor
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def valor_excel(valor):
    """Convierte lo que devuelve SQL Server a algo que Excel entienda."""
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, (date, datetime, time)):
        return valor
    return valor


def columnas_de(tabla: str) -> list[str]:
    definicion = NOMBRE_TABLA.get(tabla)
    if definicion and definicion.get("columnas_exportar"):
        return list(definicion["columnas_exportar"])
    return columnas(tabla)


def clave_primaria(tabla: str) -> str:
    return "Id_" + {
        "CLIENTE": "Cliente",
        "PODOLOGO": "Podologo",
        "RECEPCIONISTA": "Recep",
        "SERVICIO": "Servicio",
        "MEDICAMENTO": "Medicamento",
        "CITA": "Cita",
        "ATENCION": "Atencion",
        "DETALLE_ATENCION": "Detalle_Aten",
        "DETALLE_MEDICAMENTO": "Detalle_Med",
        "BOLETA": "Boleta",
        "PAGO": "Pago",
        "USUARIO": "Usuario",
        "AUDITORIA": "Auditoria",
    }[tabla]


def leer_tabla(tabla: str) -> list[dict]:
    return consultar("SELECT * FROM %s" % tabla)


# ---------------------------------------------------------------------------
#  Exportar
# ---------------------------------------------------------------------------
def exportar_json() -> dict:
    tablas = {}
    for definicion in TABLAS:
        nombre = definicion["nombre"]
        filas = leer_tabla(nombre)
        tablas[nombre] = [{c: valor_excel(f.get(c)) for c in columnas_de(nombre)} for f in filas]
    return {
        "sistema": "Clínica Podológica",
        "base_de_datos": BASE_DATOS,
        "generado": datetime.now().isoformat(timespec="seconds"),
        "tablas": tablas,
    }


def exportar_excel() -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA_INSTRUCCIONES
    escribir_instrucciones(hoja)

    for definicion in TABLAS:
        hoja = libro.create_sheet(definicion["nombre"])
        escribir_hoja(hoja, definicion)

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def escribir_hoja(hoja, definicion: dict) -> None:
    columnas_tabla = columnas_de(definicion["nombre"])
    hoja.append(columnas_tabla)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0B5F6B")
        celda.alignment = Alignment(horizontal="center")
    for fila in leer_tabla(definicion["nombre"]):
        hoja.append([valor_excel(fila.get(c)) for c in columnas_tabla])
    hoja.freeze_panes = "A2"
    for indice, nombre in enumerate(columnas_tabla, 1):
        hoja.column_dimensions[get_column_letter(indice)].width = max(12, min(34, len(nombre) + 4))


def escribir_instrucciones(hoja) -> None:
    titulo = Font(bold=True, size=14, color="0B5F6B")
    subtitulo = Font(bold=True, size=11)
    hoja.column_dimensions["A"].width = 26
    hoja.column_dimensions["B"].width = 46
    hoja.column_dimensions["C"].width = 78

    def linea(a="", b="", c="", estilo=None):
        hoja.append([a, b, c])
        if estilo:
            hoja.cell(row=hoja.max_row, column=1).font = estilo

    linea("CLÍNICA PODOLÓGICA — RESPALDO DE DATOS", estilo=titulo)
    linea("Base de datos", BASE_DATOS, "Generado el %s" % datetime.now().strftime("%d/%m/%Y %H:%M"))
    linea()
    linea("CÓMO USAR ESTE ARCHIVO", estilo=subtitulo)
    linea("Exportar", "Cada hoja de este libro es una tabla de la base de datos.",
          "El nombre de la hoja es el nombre exacto de la tabla. No lo cambie.")
    linea("Importar", "Suba este archivo desde el módulo Copia de seguridad.",
          "Se agrega lo que falta y se actualiza lo que coincide por su clave natural. Nunca se borra nada.")
    linea("Columnas", "La primera fila de cada hoja son los nombres de las columnas.",
          "Puede reordenar las columnas y dejar celdas vacías; las que no incluya se mantienen como están.")
    linea("Filas vacías", "Las filas sin datos se ignoran.",
          "Puede borrar filas y agregar nuevas libremente.")
    linea()
    linea("ORDEN DE LAS TABLAS", estilo=subtitulo)
    linea("", "El orden de las hojas respeta las claves foráneas:",
          "primero Cliente, Podólogo y Recepcionista; después Citas, Atenciones, Boletas y Pagos.")
    linea()
    linea("HOJAS DE ESTE LIBRO", estilo=subtitulo)
    linea("Hoja", "Qué contiene", "Restricciones y reglas")
    for definicion in TABLAS:
        nombre = definicion["nombre"]
        aviso = "" if definicion.get("importar", True) else "  (NO se importa)"
        linea(nombre + aviso, definicion["detalle"], definicion["restricciones"])
    linea()
    linea("REGLAS DEL NEGOCIO", estilo=subtitulo)
    linea("R1", "Identidad única de la cita",
          "Un podólogo no puede tener dos citas activas a la misma fecha y hora.")
    linea("R2", "Consistencia de pago",
          "El importe de la boleta es la suma de sus consumos y los pagos no pueden superarlo. "
          "La boleta pasa a Pagada solo cuando los pagos igualan el importe.")
    linea("R3", "Seguridad de roles",
          "Solo el Administrador ve la auditoría y gestiona usuarios.")
    linea("R4", "Inmutabilidad de registros",
          "Una boleta pagada o una atención ya facturada no se pueden modificar ni eliminar.")
    linea("R5", "Referencialidad obligatoria",
          "No puede existir una boleta sin una atención y un cliente registrados.")
    linea()
    linea("SEGURIDAD", estilo=subtitulo)
    linea("Contraseñas", "No se exportan nunca.",
          "La tabla USUARIO se incluye solo como referencia; en la base se guardan hashes PBKDF2-SHA256.")


# ---------------------------------------------------------------------------
#  Importar
# ---------------------------------------------------------------------------
class Resumen:
    def __init__(self):
        self.insertados: dict[str, int] = {}
        self.actualizados: dict[str, int] = {}
        self.omitidos: dict[str, int] = {}
        self.avisos: list[str] = []

    def sumar(self, grupo: dict, tabla: str, cantidad: int = 1) -> None:
        grupo[tabla] = grupo.get(tabla, 0) + cantidad

    def aviso(self, texto: str) -> None:
        if texto not in self.avisos and len(self.avisos) < 25:
            self.avisos.append(texto)

    def como_dict(self) -> dict:
        return {
            "insertados": self.insertados,
            "actualizados": self.actualizados,
            "omitidos": self.omitidos,
            "avisos": self.avisos,
            "total_insertados": sum(self.insertados.values()),
            "total_actualizados": sum(self.actualizados.values()),
            "total_omitidos": sum(self.omitidos.values()),
        }


def importar_json(paquete: dict) -> dict:
    if not isinstance(paquete, dict):
        raise ValueError("El archivo JSON no tiene el formato esperado.")
    tablas = paquete.get("tablas") if isinstance(paquete.get("tablas"), dict) else paquete
    limpias = {}
    for nombre, filas in tablas.items():
        if nombre in NOMBRE_TABLA and isinstance(filas, list):
            limpias[nombre] = [f for f in filas if isinstance(f, dict)]
    if not limpias:
        raise ValueError("El archivo JSON no contiene ninguna tabla reconocida.")
    return _importar(limpias)


def importar_excel(contenido: bytes) -> dict:
    try:
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception:
        raise ValueError("No se pudo leer el archivo. Asegúrese de que sea un Excel (.xlsx) válido.")

    limpias = {}
    for definicion in TABLAS:
        nombre = definicion["nombre"]
        hoja = libro[nombre] if nombre in libro.sheetnames else None
        if hoja is None:
            for posible in libro.sheetnames:
                if posible.strip().upper() in (nombre.upper(), definicion["etiqueta"].upper()):
                    hoja = libro[posible]
                    break
        if hoja is None:
            continue
        filas = leer_hoja(hoja)
        if filas:
            limpias[nombre] = filas
    libro.close()

    if not limpias:
        raise ValueError("El archivo no tiene ninguna hoja con el nombre de una tabla "
                         "(se esperaba CLIENTE, PODOLOGO, RECEPCIONISTA, ...).")
    return _importar(limpias)


def leer_hoja(hoja) -> list[dict]:
    filas = []
    cabeceras: list[str] = []
    for indice, fila in enumerate(hoja.iter_rows(values_only=True)):
        if indice == 0:
            cabeceras = [a_texto(c) for c in fila]
            continue
        if fila is None or all(c is None or a_texto(c) == "" for c in fila):
            continue
        registro = {}
        for posicion, valor in enumerate(fila):
            if posicion < len(cabeceras) and cabeceras[posicion]:
                registro[cabeceras[posicion]] = valor
        if registro:
            filas.append(registro)
    return filas


def _importar(tablas: dict) -> dict:
    resumen = Resumen()
    mapa: dict[tuple, int] = {}
    tx = Tx()
    try:
        _padres(tx, tablas, mapa, resumen)
        _citas(tx, tablas, mapa, resumen)
        _atenciones(tx, tablas, mapa, resumen)
        _detalles(tx, tablas, mapa, resumen)
        _boletas(tx, tablas, mapa, resumen)
        _pagos(tx, tablas, mapa, resumen)
        tx.confirmar()
    except Exception:
        tx.cn.rollback()
        raise
    finally:
        tx.cerrar()

    for definicion in TABLAS:
        if not definicion.get("importar", True) and definicion["nombre"] in tablas:
            resumen.aviso("La hoja %s no se importa (%s)." % (definicion["nombre"], definicion["restricciones"][:60]))
    return resumen.como_dict()


def resolver(mapa: dict, tabla: str, id_viejo) -> int | None:
    if id_viejo is None:
        return None
    try:
        return mapa.get((tabla, int(id_viejo)))
    except (TypeError, ValueError):
        return None


def _insertar(tx, tabla: str, fila: dict, permite_clave: bool = False) -> int:
    permitidas = [c for c in columnas_de(tabla) if c in fila]
    if not permite_clave:
        pk = clave_primaria(tabla)
        permitidas = [c for c in permitidas if c != pk]
    if not permitidas:
        return 0
    sql = "INSERT INTO %s (%s) VALUES (%s)" % (
        tabla, ", ".join(permitidas), ", ".join("?" * len(permitidas)))
    return tx.insertar(sql, tuple(a_valor(fila[c]) for c in permitidas))


def _actualizar(tx, tabla: str, pk_valor, fila: dict) -> None:
    pk = clave_primaria(tabla)
    campos = [c for c in columnas_de(tabla) if c in fila and c != pk]
    if not campos:
        return
    sql = "UPDATE %s SET %s WHERE %s = ?" % (
        tabla, ", ".join(c + " = ?" for c in campos), pk)
    tx.ejecutar(sql, tuple(a_valor(fila[c]) for c in campos) + (pk_valor,))


def _buscar_por_clave(tx, tabla: str, fila: dict):
    """Busca un registro existente usando su clave natural."""
    for columna in CLAVES.get(tabla, ()):
        valor = a_texto(fila.get(columna))
        if valor:
            encontrado = tx.uno(
                "SELECT %s FROM %s WHERE %s = ?" % (clave_primaria(tabla), tabla, columna), (valor,))
            if encontrado:
                return encontrado
    return None


def _padres(tx, tablas: dict, mapa: dict, resumen: Resumen) -> None:
    for nombre in ("CLIENTE", "PODOLOGO", "RECEPCIONISTA", "SERVICIO", "MEDICAMENTO"):
        definicion = NOMBRE_TABLA[nombre]
        if not definicion.get("importar", True):
            continue
        for fila in tablas.get(nombre) or []:
            obligatorio = {"CLIENTE": "Noms_Cliente", "PODOLOGO": "Nom_Podologo",
                           "RECEPCIONISTA": "Noms_Recep", "SERVICIO": "Nom_serv",
                           "MEDICAMENTO": "Nom_Med"}[nombre]
            if not a_texto(fila.get(obligatorio)):
                resumen.sumar(resumen.omitidos, nombre)
                resumen.aviso("Se omitió una fila de %s sin %s." % (nombre, obligatorio))
                continue
            existe = _buscar_por_clave(tx, nombre, fila)
            if existe:
                nuevo = existe[clave_primaria(nombre)]
                _actualizar(tx, nombre, nuevo, fila)
                resumen.sumar(resumen.actualizados, nombre)
            else:
                nuevo = _insertar(tx, nombre, fila)
                resumen.sumar(resumen.insertados, nombre)
            if fila.get(clave_primaria(nombre)) is not None:
                mapa[(nombre, int(fila[clave_primaria(nombre)]))] = nuevo


def _citas(tx, tablas: dict, mapa: dict, resumen: Resumen) -> None:
    for fila in tablas.get("CITA") or []:
        id_cliente = resolver(mapa, "CLIENTE", fila.get("Id_Cliente"))
        id_podologo = resolver(mapa, "PODOLOGO", fila.get("Id_Podologo"))
        id_recep = resolver(mapa, "RECEPCIONISTA", fila.get("Id_Recep"))
        if not id_cliente or not id_podologo:
            resumen.sumar(resumen.omitidos, "CITA")
            resumen.aviso("Se omitió una cita porque su cliente o podólogo no venía en el archivo.")
            continue
        datos = dict(fila, Id_Cliente=id_cliente, Id_Podologo=id_podologo, Id_Recep=id_recep)
        fecha = a_valor(fila.get("Fe_Cita"))
        hora = a_valor(fila.get("Hora_Cita"))
        existe = None
        if fecha is not None and hora is not None:
            existe = tx.uno(
                """SELECT Id_Cita FROM CITA WHERE Id_Cliente = ? AND Id_Podologo = ?
                   AND Fe_Cita = ? AND Hora_Cita = ?""",
                (id_cliente, id_podologo, fecha, hora))
        if existe:
            nuevo = existe["Id_Cita"]
            _actualizar(tx, "CITA", nuevo, datos)
            resumen.sumar(resumen.actualizados, "CITA")
        else:
            nuevo = _insertar(tx, "CITA", datos)
            resumen.sumar(resumen.insertados, "CITA")
        if fila.get("Id_Cita") is not None:
            mapa[("CITA", int(fila["Id_Cita"]))] = nuevo


def _atenciones(tx, tablas: dict, mapa: dict, resumen: Resumen) -> None:
    for fila in tablas.get("ATENCION") or []:
        id_cita = resolver(mapa, "CITA", fila.get("Id_Cita"))
        if not id_cita:
            resumen.sumar(resumen.omitidos, "ATENCION")
            resumen.aviso("Se omitió una atención porque su cita no venía en el archivo.")
            continue
        datos = dict(fila, Id_Cita=id_cita)
        existe = tx.uno("SELECT Id_Atencion FROM ATENCION WHERE Id_Cita = ?", (id_cita,))
        if existe:
            nuevo = existe["Id_Atencion"]
            _actualizar(tx, "ATENCION", nuevo, datos)
            resumen.sumar(resumen.actualizados, "ATENCION")
        else:
            nuevo = _insertar(tx, "ATENCION", datos)
            resumen.sumar(resumen.insertados, "ATENCION")
        if fila.get("Id_Atencion") is not None:
            mapa[("ATENCION", int(fila["Id_Atencion"]))] = nuevo


def _detalles(tx, tablas: dict, mapa: dict, resumen: Resumen) -> None:
    for nombre, tabla_hija in (("DETALLE_ATENCION", "SERVICIO"), ("DETALLE_MEDICAMENTO", "MEDICAMENTO")):
        columna_hija = "Id_Servicio" if tabla_hija == "SERVICIO" else "Id_Medicamento"
        columna_fk = "Id_Detalle_Aten" if tabla_hija == "SERVICIO" else "Id_Detalle_Med"
        for fila in tablas.get(nombre) or []:
            id_atencion = resolver(mapa, "ATENCION", fila.get("Id_Atencion"))
            id_hijo = resolver(mapa, tabla_hija, fila.get(columna_hija))
            if not id_atencion or not id_hijo:
                resumen.sumar(resumen.omitidos, nombre)
                resumen.aviso("Se omitió una fila de %s porque su atención o %s no venía en el archivo."
                              % (nombre, tabla_hija.lower()))
                continue
            datos = dict(fila, Id_Atencion=id_atencion)
            datos[columna_hija] = id_hijo
            existe = tx.uno(
                "SELECT %s FROM %s WHERE Id_Atencion = ? AND %s = ?" % (columna_fk, nombre, columna_hija),
                (id_atencion, id_hijo))
            if existe:
                _actualizar(tx, nombre, existe[columna_fk], datos)
                resumen.sumar(resumen.actualizados, nombre)
            else:
                _insertar(tx, nombre, datos)
                resumen.sumar(resumen.insertados, nombre)


def _boletas(tx, tablas: dict, mapa: dict, resumen: Resumen) -> None:
    for fila in tablas.get("BOLETA") or []:
        id_atencion = resolver(mapa, "ATENCION", fila.get("Id_Atencion"))
        if not id_atencion:
            resumen.sumar(resumen.omitidos, "BOLETA")
            resumen.aviso("Se omitió una boleta porque su atención no venía en el archivo.")
            continue
        id_recep = resolver(mapa, "RECEPCIONISTA", fila.get("Id_Recep"))
        datos = dict(fila, Id_Atencion=id_atencion, Id_Recep=id_recep)
        existe = tx.uno("SELECT Id_Boleta, Estado_Pago FROM BOLETA WHERE Id_Atencion = ?", (id_atencion,))
        if existe:
            if existe["Estado_Pago"] == "Pagada":
                resumen.sumar(resumen.omitidos, "BOLETA")
                resumen.aviso("La boleta de la atención %d ya está pagada: regla R4, no se modificó." % id_atencion)
                mapa[("BOLETA", int(fila.get("Id_Boleta") or 0))] = existe["Id_Boleta"]
                continue
            nuevo = existe["Id_Boleta"]
            _actualizar(tx, "BOLETA", nuevo, datos)
            resumen.sumar(resumen.actualizados, "BOLETA")
        else:
            nuevo = _insertar(tx, "BOLETA", datos)
            resumen.sumar(resumen.insertados, "BOLETA")
        if fila.get("Id_Boleta") is not None:
            mapa[("BOLETA", int(fila["Id_Boleta"]))] = nuevo


def _pagos(tx, tablas: dict, mapa: dict, resumen: Resumen) -> None:
    for fila in tablas.get("PAGO") or []:
        id_boleta = resolver(mapa, "BOLETA", fila.get("Id_Boleta"))
        if not id_boleta:
            resumen.sumar(resumen.omitidos, "PAGO")
            resumen.aviso("Se omitió un pago porque su boleta no venía en el archivo.")
            continue
        datos = dict(fila, Id_Boleta=id_boleta)
        existe = tx.uno(
            """SELECT Id_Pago FROM PAGO WHERE Id_Boleta = ? AND Tipo_Pago = ?
               AND Monto = ? AND Fe_Pago = ?""",
            (id_boleta, a_texto(fila.get("Tipo_Pago")), a_valor(fila.get("Monto")),
             a_valor(fila.get("Fe_Pago"))))
        if existe:
            resumen.sumar(resumen.omitidos, "PAGO")
            continue
        _insertar(tx, "PAGO", datos)
        resumen.sumar(resumen.insertados, "PAGO")
        _recalcular_boleta(tx, id_boleta)


def _recalcular_boleta(tx, id_boleta: int) -> None:
    """Deja el Estado_Pago acorde a los pagos registrados (regla R2)."""
    pagado = tx.uno("SELECT ISNULL(SUM(Monto),0) t FROM PAGO WHERE Id_Boleta = ?", (id_boleta,))["t"]
    importe = tx.uno("SELECT Importe FROM BOLETA WHERE Id_Boleta = ?", (id_boleta,))["Importe"]
    estado = "Pagada" if float(pagado) >= float(importe) else "Pendiente"
    tx.ejecutar("UPDATE BOLETA SET Estado_Pago = ? WHERE Id_Boleta = ? AND Estado_Pago <> 'Pagada'",
                (estado, id_boleta))


def contar_tablas() -> list[dict]:
    """Cuenta los registros de cada tabla, para mostrarlo en pantalla."""
    resultado = []
    for definicion in TABLAS:
        resultado.append({
            "tabla": definicion["nombre"],
            "etiqueta": definicion["etiqueta"],
            "filas": escalar("SELECT COUNT(*) FROM %s" % definicion["nombre"], (), 0),
        })
    return resultado
